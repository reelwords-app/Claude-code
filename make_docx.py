"""
Convert paper_pinn_polymer_flood.md to a formatted Word (.docx) document.
Uses python-docx; handles headings, paragraphs, tables, block-quotes, bold.
"""
import re
from docx import Document
from docx.shared import Pt, RGBColor, Inches, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

MD_PATH  = '/home/user/Claude-code/paper_pinn_polymer_flood.md'
OUT_PATH = '/home/user/Claude-code/paper_pinn_polymer_flood.docx'

doc = Document()

# ── Page margins ──────────────────────────────────────────────────────────────
section = doc.sections[0]
section.page_width  = Inches(8.5)
section.page_height = Inches(11)
section.left_margin = section.right_margin = Inches(1.0)
section.top_margin  = section.bottom_margin = Inches(1.0)

# ── Default font ──────────────────────────────────────────────────────────────
style = doc.styles['Normal']
style.font.name = 'Times New Roman'
style.font.size = Pt(11)

def set_heading(level, text):
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        run.font.name = 'Times New Roman'
        run.font.color.rgb = RGBColor(0, 0, 0)
    return h

def add_para(text, style_name='Normal', bold=False, italic=False,
             align=WD_ALIGN_PARAGRAPH.LEFT, size=None):
    p = doc.add_paragraph(style=style_name)
    p.alignment = align
    # Handle inline bold/italic markdown: **text** and *text*
    segments = re.split(r'(\*\*.*?\*\*|\*.*?\*)', text)
    for seg in segments:
        if seg.startswith('**') and seg.endswith('**'):
            r = p.add_run(seg[2:-2])
            r.bold = True
        elif seg.startswith('*') and seg.endswith('*'):
            r = p.add_run(seg[1:-1])
            r.italic = True
        else:
            r = p.add_run(seg)
        r.font.name = 'Times New Roman'
        r.font.size = Pt(size or 11)
        if bold:   r.bold   = True
        if italic: r.italic = True
    return p

def add_quote(text):
    """Block-quote style for figure captions."""
    p = doc.add_paragraph(style='Normal')
    p.paragraph_format.left_indent  = Inches(0.4)
    p.paragraph_format.right_indent = Inches(0.4)
    # Strip leading > and whitespace
    text = re.sub(r'^>\s*', '', text).strip()
    segments = re.split(r'(\*\*.*?\*\*|\*.*?\*)', text)
    for seg in segments:
        if seg.startswith('**') and seg.endswith('**'):
            r = p.add_run(seg[2:-2]); r.bold = True
        elif seg.startswith('*') and seg.endswith('*'):
            r = p.add_run(seg[1:-1]); r.italic = True
        else:
            r = p.add_run(seg)
        r.font.name = 'Times New Roman'
        r.font.size = Pt(10)
        r.font.color.rgb = RGBColor(60, 60, 60)
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after  = Pt(4)
    # Add left border
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    left = OxmlElement('w:left')
    left.set(qn('w:val'), 'single')
    left.set(qn('w:sz'), '6')
    left.set(qn('w:space'), '12')
    left.set(qn('w:color'), '2980B9')
    pBdr.append(left)
    pPr.append(pBdr)
    return p

def add_table_from_md(lines):
    """Parse a markdown table and add it as a docx table."""
    rows = []
    for ln in lines:
        if re.match(r'\s*\|[-| :]+\|\s*$', ln):
            continue  # separator row
        cells = [c.strip() for c in ln.strip().strip('|').split('|')]
        rows.append(cells)
    if not rows:
        return

    n_cols = max(len(r) for r in rows)
    tbl = doc.add_table(rows=len(rows), cols=n_cols)
    tbl.style = 'Table Grid'
    for ri, row_data in enumerate(rows):
        for ci, cell_text in enumerate(row_data):
            if ci >= n_cols:
                break
            cell = tbl.cell(ri, ci)
            cell.text = ''
            p = cell.paragraphs[0]
            # Remove **bold** markers
            clean = re.sub(r'\*\*(.*?)\*\*', r'\1', cell_text)
            clean = re.sub(r'\*(.*?)\*',    r'\1', clean)
            run = p.add_run(clean)
            run.font.name = 'Times New Roman'
            run.font.size = Pt(10)
            if ri == 0:
                run.bold = True
    doc.add_paragraph()

# ── Parse and render markdown ─────────────────────────────────────────────────
with open(MD_PATH, encoding='utf-8') as f:
    lines = f.readlines()

i = 0
while i < len(lines):
    ln = lines[i].rstrip('\n')

    # Headings
    if ln.startswith('#### '):
        set_heading(4, ln[5:].strip()); i += 1; continue
    if ln.startswith('### '):
        set_heading(3, ln[4:].strip()); i += 1; continue
    if ln.startswith('## '):
        set_heading(2, ln[3:].strip()); i += 1; continue
    if ln.startswith('# '):
        p = set_heading(1, ln[2:].strip())
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        i += 1; continue

    # Horizontal rule
    if re.match(r'^---+\s*$', ln):
        doc.add_paragraph(); i += 1; continue

    # Block quote (figure captions)
    if ln.startswith('>'):
        add_quote(ln); i += 1; continue

    # Markdown table: collect all consecutive table lines
    if ln.startswith('|'):
        tbl_lines = []
        while i < len(lines) and lines[i].startswith('|'):
            tbl_lines.append(lines[i].rstrip('\n'))
            i += 1
        add_table_from_md(tbl_lines)
        continue

    # Numbered list items
    m = re.match(r'^(\d+)\.\s+(.*)', ln)
    if m:
        p = doc.add_paragraph(style='List Number')
        text = m.group(2)
        segments = re.split(r'(\*\*.*?\*\*|\*.*?\*)', text)
        for seg in segments:
            if seg.startswith('**') and seg.endswith('**'):
                r = p.add_run(seg[2:-2]); r.bold = True
            elif seg.startswith('*') and seg.endswith('*'):
                r = p.add_run(seg[1:-1]); r.italic = True
            else:
                r = p.add_run(seg)
            r.font.name = 'Times New Roman'; r.font.size = Pt(11)
        i += 1; continue

    # Bullet list items
    if ln.startswith('- ') or ln.startswith('* '):
        p = doc.add_paragraph(style='List Bullet')
        text = ln[2:]
        segments = re.split(r'(\*\*.*?\*\*|\*.*?\*)', text)
        for seg in segments:
            if seg.startswith('**') and seg.endswith('**'):
                r = p.add_run(seg[2:-2]); r.bold = True
            elif seg.startswith('*') and seg.endswith('*'):
                r = p.add_run(seg[1:-1]); r.italic = True
            else:
                r = p.add_run(seg)
            r.font.name = 'Times New Roman'; r.font.size = Pt(11)
        i += 1; continue

    # Empty line
    if ln.strip() == '':
        i += 1; continue

    # Regular paragraph
    add_para(ln.strip()); i += 1

doc.save(OUT_PATH)
print(f'Word document written: {OUT_PATH}')
